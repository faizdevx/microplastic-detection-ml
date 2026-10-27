"""Reviewer_Response.docx claims against the evidence in this repository.

    python tools/claims_vs_evidence.py            write docs/claims_vs_evidence.md
    python tools/claims_vs_evidence.py --check    also exit 1 if a `supported` claim has no matching number

The docx is opened read-only and never modified. Mapping: docs/claims_map.yaml.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Reviewer_Response.docx"
MAP = ROOT / "docs" / "claims_map.yaml"
OUT = ROOT / "docs" / "claims_vs_evidence.md"
STATUSES = ["supported", "contradicted", "not reproducible here", "different protocol", "consistent"]
KEYWORDS = re.compile(
    r"mAP|SVR|Lasso|Random Forest|hold-out|final_dataset|5-fold|fold|leakage|partition|"
    r"mean_particle_area|green_std|max_particle_area|repository|stratified|F1|precision|recall|MAE|"
    r"Classification Stage|Quantification Stage|feature-based",
    re.I)
HEADING = re.compile(
    r"^(Clarification|Prevention|Dataset and|Hyperparameters|Classification Performance)", re.I)


# ------------------------------------------------------------------- docx ---------------
def read_docx(path: Path) -> tuple[list[dict], list[list[list[str]]]]:
    import docx

    d = docx.Document(str(path))
    reviewer, comment, in_response = "", "", False
    claims = []
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if re.fullmatch(r"Reviewer \d+", t):
            reviewer, comment, in_response = t, "", False
            continue
        m = re.match(r"comment\s+(\d+)", t, re.I)
        if m and len(t) < 20:
            comment, in_response = f"Comment {m.group(1)}", False
            continue
        if t.startswith("Response"):
            in_response = True
            if len(t) < 14:
                continue
        if HEADING.match(t) or (t.endswith(":") and not re.search(r"\d", t)):
            continue  # section headings and lead-in lines are not claims
        if in_response and (re.search(r"\d", t) or KEYWORDS.search(t)):
            claims.append({"id": f"p{i}", "where": f"{reviewer}, {comment}".strip(", "), "text": t})
    tables, seen = [], set()
    for ti, tb in enumerate(d.tables):
        rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in tb.rows]
        key = json.dumps(rows)
        if key in seen:
            continue
        seen.add(key)
        tables.append(rows)
        for ri, r in enumerate(rows[1:], 1):
            claims.append({"id": f"t{len(tables)}r{ri}", "where": "regression table (Comments 2 and 13)",
                           "text": f"{r[0]}: R2 {r[1]}, MAE {r[2]}" if len(r) >= 3 else " | ".join(r)})
    return claims, tables


# ---------------------------------------------------------------------- evidence --------
def resolve_json(path: str, key: str):
    obj = json.loads((ROOT / path).read_text(encoding="utf8"))
    for part in re.findall(r"[^.\[\]]+|\[\d+\]", key):
        obj = obj[int(part[1:-1])] if part.startswith("[") else obj[part]
    return obj


def resolve(ref: str):
    if ref.startswith("csv_rows:"):
        with open(ROOT / ref.split(":", 1)[1], encoding="utf8", newline="") as fh:
            return sum(1 for _ in csv.DictReader(fh))
    if ref.startswith("csv_class:"):
        _, path, col, val = ref.split(":", 3)
        with open(ROOT / path, encoding="utf8", newline="") as fh:
            return sum(1 for r in csv.DictReader(fh) if r[col] == val)
    if ref.startswith("grep:"):
        path, pat = ref[5:].split("::", 1)
        return len(re.findall(pat, (ROOT / path).read_text(encoding="utf8", errors="replace")))
    path, key = ref.split("::", 1)
    return resolve_json(path, key)


def flat_numbers(v) -> list[float]:
    if isinstance(v, bool):
        return []
    if isinstance(v, (int, float)):
        return [float(v)]
    if isinstance(v, list):
        return [x for e in v for x in flat_numbers(e)]
    return []


def matches(claim_number: float, values: list[float], decimals: int) -> bool:
    """Equal at the claim's printed precision; a percentage may match a fraction."""
    for v in values:
        for cand in (v, v * 100.0):
            if round(cand, decimals) == round(claim_number, decimals):
                return True
    return False


def decimals_of(s: float) -> int:
    t = repr(float(s))
    return len(t.split(".")[1]) if "." in t and t.split(".")[1] != "0" else 0


def fmt_value(v) -> str:
    if isinstance(v, float):
        return f"{v:.4g}" if abs(v) < 1 else f"{v:.2f}"
    return json.dumps(v) if isinstance(v, (list, dict)) else str(v)


def check_entry(cid: str, entry: dict) -> tuple[list[str], list[str]]:
    """Return (evidence lines, problems)."""
    lines, problems, vals = [], [], []
    for ev in entry.get("evidence", []):
        try:
            v = resolve(ev["ref"])
        except Exception as exc:
            problems.append(f"{cid}: cannot resolve {ev['ref']} ({type(exc).__name__})")
            continue
        vals += flat_numbers(v)
        lines.append(f"{ev.get('label', ev['ref'])} = {fmt_value(v)} (`{ev['ref'].split('::')[0]}`)")
    if entry["status"] == "supported":
        nums = entry.get("claim_numbers")
        if not nums:
            problems.append(f"{cid}: supported without claim_numbers")
        for n in nums or []:
            if not matches(float(n), vals, decimals_of(n)):
                problems.append(f"{cid}: supported but {n} matches no evidence value")
    if entry["status"] not in STATUSES:
        problems.append(f"{cid}: unknown status {entry['status']!r}")
    return lines, problems


# --------------------------------------------------------------------------- main -------
def build() -> tuple[str, list[str]]:
    before = hashlib.sha256(DOCX.read_bytes()).hexdigest()
    claims, _tables = read_docx(DOCX)
    mp = yaml.safe_load(MAP.read_text(encoding="utf8"))
    facts = {}
    problems: list[str] = []
    for k, f in mp.get("facts", {}).items():
        try:
            v = resolve(f["ref"])
            facts[k] = f"{v:.3f}" if isinstance(v, float) else v
        except Exception as exc:
            problems.append(f"fact {k}: {type(exc).__name__}")
    rows, counts = [], Counter()
    for c in claims:
        entry = mp["claims"].get(c["id"])
        if entry is None:
            rows.append((c, "unmapped", "no evidence mapped", ""))
            counts["unmapped"] += 1
            continue
        ev, pr = check_entry(c["id"], entry)
        problems += pr
        counts[entry["status"]] += 1
        rows.append((c, entry["status"], "<br>".join(ev) or "no evidence in the repository",
                     entry.get("note", "")))
    stale = sorted(set(mp["claims"]) - {c["id"] for c in claims})
    problems += [f"{s}: in claims_map.yaml but not found in the docx" for s in stale]
    L = ["# Claims in `Reviewer_Response.docx` against the repository", "",
         "Generated by `python tools/claims_vs_evidence.py` from the docx (read-only) and "
         "`docs/claims_map.yaml`. Statuses: `supported` (a number in the claim equals a value in a "
         "generated file), `contradicted`, `not reproducible here`, `different protocol` (other data or "
         "splits, so the numbers cannot be compared), `consistent` (qualitative statement that matches), "
         "`unmapped` (no evidence mapped yet).", "",
         "| Status | Claims |", "|---|---|"]
    for s in STATUSES + ["unmapped"]:
        L.append(f"| {s} | {counts.get(s, 0)} |")
    L += ["", f"{len(claims)} claims extracted.", "",
          "| Id | Where | Claim | Status | Evidence in the repository | Note |", "|---|---|---|---|---|---|"]
    for c, status, ev, note in rows:
        txt = c["text"].replace("|", "/")
        L.append(f"| {c['id']} | {c['where']} | {txt} | {status} | {ev} | {note} |")
    L += ["", "## Discrepancies to resolve", ""]
    for d in mp.get("discrepancies", []):
        L.append("- " + d.format(**facts))
    text = "\n".join(L) + "\n"
    after = hashlib.sha256(DOCX.read_bytes()).hexdigest()
    if before != after:
        raise SystemExit("the docx changed while running; this must not happen")
    return text, problems + [f"unmapped: {c['id']}" for c, s, _, _ in rows if s == "unmapped"]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf8", errors="replace")
    text, problems = build()
    OUT.write_text(text, encoding="utf8", newline="\n")
    print(f"wrote {OUT.relative_to(ROOT)}; {len(problems)} problem(s)")
    for p in problems:
        print("  ", p)
    hard = [p for p in problems if not p.startswith("unmapped")]
    return 1 if args.check and hard else 0


if __name__ == "__main__":
    sys.exit(main())
